import os
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pptx import Presentation
from pptx.util import Inches as PptxInches, Pt as PptxPt

def create_report():
    doc = Document()

    # --- Cover Page ---
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("DRIVER DROWSINESS DETECTION SYSTEM\n")
    run.bold = True
    run.font.name = 'Times New Roman'
    run.font.size = Pt(18)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.line_spacing = 1.5
    run2 = p2.add_run("PROJECT-I & II REPORT\n")
    run2.bold = True
    run2.font.name = 'Times New Roman'
    run2.font.size = Pt(14)

    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run3 = p3.add_run("Submitted by\n")
    run3.italic = True
    run3.font.name = 'Times New Roman'
    run3.font.size = Pt(14)

    p4 = doc.add_paragraph()
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run4 = p4.add_run("[YOUR NAME HERE]\n[YOUR ENROLLMENT NUMBER]\n")
    run4.bold = True
    run4.font.name = 'Times New Roman'
    run4.font.size = Pt(16)

    p5 = doc.add_paragraph()
    p5.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run5 = p5.add_run("in partial fulfillment for the award of the degree of\n")
    run5.font.name = 'Times New Roman'
    run5.font.size = Pt(14)

    p6 = doc.add_paragraph()
    p6.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p6.paragraph_format.line_spacing = 1.5
    run6 = p6.add_run("B.Tech (Computer Science & Engineering)\n")
    run6.bold = True
    run6.italic = True
    run6.font.name = 'Times New Roman'
    run6.font.size = Pt(16)

    p7 = doc.add_paragraph()
    p7.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run7 = p7.add_run("Under the supervision of\n")
    run7.italic = True
    run7.font.name = 'Times New Roman'
    run7.font.size = Pt(14)

    p8 = doc.add_paragraph()
    p8.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run8 = p8.add_run("[NAME OF THE SUPERVISOR]\n")
    run8.bold = True
    run8.font.name = 'Times New Roman'
    run8.font.size = Pt(16)

    p9 = doc.add_paragraph()
    p9.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run9 = p9.add_run("Department of Computer Science & Engineering\nSchool of Engineering Sciences & Technology\n")
    run9.font.name = 'Times New Roman'
    run9.font.size = Pt(16)

    p10 = doc.add_paragraph()
    p10.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run10 = p10.add_run("JAMIA HAMDARD\n")
    run10.bold = True
    run10.font.name = 'Times New Roman'
    run10.font.size = Pt(22)

    p11 = doc.add_paragraph()
    p11.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run11 = p11.add_run("New Delhi-110062\n")
    run11.font.name = 'Times New Roman'
    run11.font.size = Pt(12)

    p12 = doc.add_paragraph()
    p12.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run12 = p12.add_run("2026\n")
    run12.bold = True
    run12.font.name = 'Times New Roman'
    run12.font.size = Pt(22)

    doc.add_page_break()

    # --- Declaration ---
    doc.add_heading('DECLARATION', level=1).alignment = WD_ALIGN_PARAGRAPH.CENTER
    decl_text = "I, Mr./Ms. [YOUR NAME], a student of B.Tech (CSE) (Enrolment No: [YOUR ENROLLMENT NUMBER]) hereby declare that the Project/Dissertation entitled “Driver Drowsiness Detection System” which is being submitted by me to the Department of Computer Science, Jamia Hamdard, New Delhi in partial fulfillment of the requirement for the award of the degree of B.Tech (CSE), is my original work and has not been submitted anywhere else for the award of any Degree, Diploma, Associateship, Fellowship or other similar title or recognition."
    doc.add_paragraph(decl_text)
    doc.add_paragraph("\n\n\n(Signature and Name of the Applicant)\nDate:\nPlace: New Delhi")
    doc.add_page_break()

    # --- Acknowledgements ---
    doc.add_heading('ACKNOWLEDGEMENTS', level=1).alignment = WD_ALIGN_PARAGRAPH.CENTER
    ack_text = "I would like to express my sincere gratitude to my supervisor, [NAME OF THE SUPERVISOR], for their invaluable guidance, continuous support, and immense knowledge throughout this project. Their mentorship has been instrumental in the successful completion of this work.\n\nI also extend my thanks to the Department of Computer Science & Engineering, Jamia Hamdard, for providing the necessary resources and environment to conduct this research and development."
    doc.add_paragraph(ack_text)
    doc.add_paragraph("\n\n(Signature)\n[YOUR NAME]\nDate:")
    doc.add_page_break()

    # --- Lists ---
    doc.add_heading('List of Abbreviations', level=1)
    doc.add_paragraph("EAR - Eye Aspect Ratio\nMAR - Mouth Aspect Ratio\nPERCLOS - Percentage of Eye Closure\nAI - Artificial Intelligence\nCV - Computer Vision")
    doc.add_page_break()

    doc.add_heading('List of Tables', level=1)
    doc.add_paragraph("[Insert List of Tables Here]")
    doc.add_page_break()

    doc.add_heading('List of Figures', level=1)
    doc.add_paragraph("[Insert List of Figures Here]")
    doc.add_page_break()

    # --- Main Content ---
    doc.add_heading('1. Title', level=1)
    doc.add_paragraph("Driver Drowsiness Detection System")

    doc.add_heading('2. Objective', level=1)
    doc.add_paragraph("The primary objective of this project is to develop a robust, real-time computer vision system that monitors driver fatigue and alertness. By analyzing facial landmarks, particularly eye closure (using the PERCLOS metric) and yawning frequency (with occlusion-aware heuristics), the system triggers real-time alerts to prevent accidents caused by driver drowsiness. It also provides a comprehensive web-based dashboard for live monitoring.")

    doc.add_heading('3. Introduction', level=1)
    doc.add_paragraph("Driver fatigue is a significant contributor to road accidents globally. Traditional methods of detecting drowsiness often rely on vehicle-based metrics (like steering pattern) which can be reactive rather than preventative. This project implements a non-intrusive, vision-based approach using Artificial Intelligence. By continuously capturing the driver's face, the system computes the Eye Aspect Ratio (EAR) and Mouth Aspect Ratio (MAR) to evaluate fatigue. The application is integrated with a Flask backend to serve a modern web interface, allowing remote or localized monitoring with high accuracy.")

    doc.add_heading('4. Problem Statement', level=1)
    doc.add_paragraph("Drowsy driving significantly impairs reaction times, vigilance, and decision-making, leading to fatal traffic accidents. Existing commercial solutions are often expensive, hardware-intensive, or prone to false positives (e.g., triggering alerts during normal speech or blinking). There is a critical need for an accessible, software-driven solution that accurately distinguishes between transient facial movements (talking, blinking) and genuine fatigue indicators (prolonged eye closure, yawning), while functioning reliably in varying lighting conditions.")

    doc.add_heading('5. SRS in IEEE format', level=1)
    doc.add_heading('5.1 Introduction', level=2)
    doc.add_paragraph("This Software Requirements Specification defines the requirements for the Driver Drowsiness Detection System.")
    doc.add_heading('5.2 Overall Description', level=2)
    doc.add_paragraph("The system comprises a Python-based computer vision engine using OpenCV and dlib/LBF landmarks, and a web dashboard built with Flask, HTML, CSS, and JavaScript. It captures webcam input, processes frames to calculate EAR and MAR, tracks head pose, and streams real-time status to the frontend.")
    doc.add_heading('5.3 Specific Requirements', level=2)
    doc.add_paragraph("Functional Requirements:\n1. The system shall capture real-time video feed from a webcam.\n2. The system shall calculate EAR to determine eye closure.\n3. The system shall calculate MAR to detect yawning.\n4. The system shall calculate PERCLOS over a rolling 3-second window.\n5. The system shall trigger an auditory and visual alert when fatigue thresholds are breached.\n6. The system shall provide a web-based dashboard displaying live metrics.\n\nNon-Functional Requirements:\n1. Performance: Frame processing must run at a minimum of 15 FPS.\n2. Accuracy: Must reduce false positives by dynamically calibrating to the user's baseline eye state.\n3. Usability: The web interface must follow modern design aesthetics.")

    doc.add_heading('6. Entity Relationship Diagram', level=1)
    doc.add_paragraph("As a computer-vision based real-time application, the core entities interacting are:\n- Driver (User being monitored)\n- Frame (Contains EAR, MAR, Pitch, Yaw metrics)\n- Alert (Triggered when metrics exceed thresholds)\n- Web Dashboard (Displays live metrics)\n\n[Please insert your ER Diagram image here]")

    doc.add_heading('7. Sample test cases', level=1)
    doc.add_paragraph("Test Case 1: Normal Blinking\n- Input: User blinks naturally (closure < 400ms).\n- Expected Output: EAR dips momentarily but PERCLOS remains below 30%. No alert triggered.\n\nTest Case 2: Prolonged Eye Closure (Drowsiness)\n- Input: User closes eyes for > 2 seconds.\n- Expected Output: PERCLOS exceeds 30% threshold. 'Drowsy' alert triggered.\n\nTest Case 3: Normal Talking\n- Input: User talks, causing MAR to fluctuate rapidly.\n- Expected Output: Yawn counter decays quickly. No false yawn alert triggered.\n\nTest Case 4: Covered Yawn\n- Input: User yawns but covers mouth with hand.\n- Expected Output: Occlusion detected combined with eye squint triggers 'Covered Yawn Detected' alert.")

    doc.add_heading('8. Snapshots of the different input and output screens', level=1)
    doc.add_paragraph("[Please insert screenshots of the web dashboard here, including:\n1. Active Monitoring state\n2. Calibration overlay\n3. Drowsy Alert state]")

    doc.add_heading('9. Conclusion', level=1)
    doc.add_paragraph("The Driver Drowsiness Detection System successfully demonstrates a reliable, cost-effective method for monitoring driver alertness. By combining the PERCLOS metric, dynamic user calibration, and advanced heuristics for occlusion-aware yawn detection, the system effectively mitigates false positives. The integration of a premium web dashboard provides an accessible interface for deploying this life-saving technology.")

    doc.add_heading('10. Limitation', level=1)
    doc.add_paragraph("1. Hardware Dependency: System accuracy drops in complete darkness without an IR camera.\n2. Extreme Angles: Face detection may fail if the driver turns their head more than 45 degrees.\n3. Heavy Occlusion: Thick-rimmed glasses with heavy glare or opaque face coverings can obstruct detection.")

    doc.add_heading('11. Bibliography', level=1)
    doc.add_paragraph("[1] T. Soukupova and J. Cech, \"Real-Time Eye Blink Detection using Facial Landmarks,\" 21st Computer Vision Winter Workshop, 2016.\n[2] W. Wierwille et al., \"Research on vehicle-based driver status/performance monitoring,\" NHTSA, 1994.\n[3] A. Rosebrock, \"Drowsiness detection with OpenCV,\" PyImageSearch, 2017.")

    doc.save('Project_Report.docx')
    print("Project_Report.docx generated successfully.")

def create_presentation():
    prs = Presentation()
    title_slide_layout = prs.slide_layouts[0]
    bullet_slide_layout = prs.slide_layouts[1]

    # Slide 1: Title
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "Driver Drowsiness Detection System"
    subtitle.text = "Student Name: [YOUR NAME]\nEnrollment Number: [YOUR ENROLLMENT NUMBER]\nSupervisor: [SUPERVISOR NAME]"

    # Slide 2: Objective
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    title_shape.text = "Objective"
    tf = body_shape.text_frame
    tf.text = "Develop a real-time computer vision system to monitor driver fatigue."
    p = tf.add_paragraph()
    p.text = "Analyze facial landmarks (EAR, MAR, Pitch, Yaw)."
    p = tf.add_paragraph()
    p.text = "Use PERCLOS (Percentage of Eye Closure) as the primary metric."
    p = tf.add_paragraph()
    p.text = "Trigger real-time auditory and visual alerts via a web dashboard."

    # Slide 3: Problem Statement
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    title_shape.text = "Problem Statement"
    tf = body_shape.text_frame
    tf.text = "Drowsy driving significantly impairs reaction times, leading to fatal accidents."
    p = tf.add_paragraph()
    p.text = "Existing solutions are often expensive or hardware-intensive."
    p = tf.add_paragraph()
    p.text = "Prone to false positives from normal talking or blinking."
    p = tf.add_paragraph()
    p.text = "Need for an accessible, software-driven, accurate solution."

    # Slide 4: Use Case Diagram
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    title_shape.text = "Use Case Diagram"
    tf = body_shape.text_frame
    tf.text = "Actors: Driver (User), Dashboard Interface, Alert System"
    p = tf.add_paragraph()
    p.text = "Use Cases: Start Monitoring, Calibrate Baseline, Trigger Alert, Stop Monitoring"
    p = tf.add_paragraph()
    p.text = "[Please insert Use Case Diagram Image Here]"

    # Slide 5: Class Diagram
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    title_shape.text = "Class Diagram"
    tf = body_shape.text_frame
    tf.text = "Main Classes:"
    p = tf.add_paragraph()
    p.text = "1. DrowsinessDetector (Core logic, threading)"
    p = tf.add_paragraph()
    p.text = "2. FaceAnalyzer (Landmark extraction, metrics calculation)"
    p = tf.add_paragraph()
    p.text = "3. AlertSystem (Audio/Visual trigger)"
    p = tf.add_paragraph()
    p.text = "[Please insert Class Diagram Image Here]"

    # Slide 6: Entity Relationship Diagram
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    title_shape.text = "Entity Relationship Diagram"
    tf = body_shape.text_frame
    tf.text = "Entities interacting in the application state:"
    p = tf.add_paragraph()
    p.text = "Session (ID, Start Time, Baseline EAR)"
    p = tf.add_paragraph()
    p.text = "Metrics Data (EAR, MAR, Pitch, Yaw)"
    p = tf.add_paragraph()
    p.text = "Alerts (Timestamp, Type: Yawn/Eyes/Pose)"
    p = tf.add_paragraph()
    p.text = "[Please insert ER Diagram Image Here]"

    # Slide 7: Gantt Chart
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    title_shape.text = "Gantt Chart (Timeline)"
    tf = body_shape.text_frame
    tf.text = "Phase 1: Literature Review & Planning (Weeks 1-2)"
    p = tf.add_paragraph()
    p.text = "Phase 2: Core CV Engine Implementation (Weeks 3-5)"
    p = tf.add_paragraph()
    p.text = "Phase 3: Web Dashboard & Integration (Weeks 6-8)"
    p = tf.add_paragraph()
    p.text = "Phase 4: Testing & Calibration (Weeks 9-10)"
    p = tf.add_paragraph()
    p.text = "[Please insert Gantt Chart Image Here]"

    # Slide 8: Conclusion
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    title_shape.text = "Conclusion"
    tf = body_shape.text_frame
    tf.text = "Successfully built a robust, real-time driver monitoring system."
    p = tf.add_paragraph()
    p.text = "Utilized PERCLOS and dynamic EAR calibration to reduce false positives."
    p = tf.add_paragraph()
    p.text = "Implemented occlusion-aware yawn detection."
    p = tf.add_paragraph()
    p.text = "Delivered a modern, 'Light Luxury' web interface for accessibility."

    # Slide 9: Bibliography
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    title_shape.text = "Bibliography"
    tf = body_shape.text_frame
    tf.text = "[1] T. Soukupova and J. Cech, 'Real-Time Eye Blink Detection using Facial Landmarks,' 21st Computer Vision Winter Workshop, 2016."
    p = tf.add_paragraph()
    p.text = "[2] W. Wierwille et al., 'Research on vehicle-based driver status/performance monitoring,' NHTSA, 1994."
    p = tf.add_paragraph()
    p.text = "[3] A. Rosebrock, 'Drowsiness detection with OpenCV,' PyImageSearch, 2017."

    prs.save('Project_Presentation.pptx')
    print("Project_Presentation.pptx generated successfully.")

if __name__ == '__main__':
    create_report()
    create_presentation()
